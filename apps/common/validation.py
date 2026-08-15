"""
Shared file-validation helpers for PDF-processing features.

Every feature that accepts uploaded PDFs (merge, split, organize, compress,
...) needs the same baseline checks, so they live here once instead of being
copy-pasted per app. Never trust a client-supplied filename or Content-Type
header alone - this validates the actual file signature and opens it with
PyMuPDF to catch corrupted/malformed/encrypted files before any processing
is attempted.
"""

import fitz
from PIL import Image
from rest_framework import serializers

PDF_MAGIC_BYTES = b"%PDF-"

DEFAULT_MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB
DEFAULT_MAX_PAGES = 2000

DEFAULT_MAX_IMAGE_SIZE = 25 * 1024 * 1024  # 25 MB
ALLOWED_IMAGE_FORMATS = {"JPEG", "PNG"}
ALLOWED_IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png")

DEFAULT_MAX_OFFICE_FILE_SIZE = 50 * 1024 * 1024  # 50 MB
# .docx/.xlsx/.pptx are all OOXML - a ZIP archive under the hood.
ZIP_MAGIC_BYTES = b"PK\x03\x04"


def validate_pdf_file(value, max_size=DEFAULT_MAX_FILE_SIZE, max_pages=DEFAULT_MAX_PAGES):
    """
    Validate an uploaded file is a real, parseable, size/page-bounded PDF.

    Raises rest_framework.serializers.ValidationError with a clear, per-file
    message on any failure. Returns the PDF's page count on success.
    """
    if value.size > max_size:
        raise serializers.ValidationError(
            f"'{value.name}' exceeds the maximum allowed size of "
            f"{max_size // (1024 * 1024)} MB."
        )

    if not value.name.lower().endswith(".pdf"):
        raise serializers.ValidationError(f"'{value.name}' is not a PDF file.")

    value.seek(0)
    header = value.read(len(PDF_MAGIC_BYTES))
    value.seek(0)

    if header != PDF_MAGIC_BYTES:
        raise serializers.ValidationError(
            f"'{value.name}' does not look like a valid PDF file."
        )

    try:
        data = value.read()
        value.seek(0)
        document = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise serializers.ValidationError(
            f"'{value.name}' could not be opened - it may be corrupted."
        )

    try:
        if document.is_encrypted and not document.authenticate(""):
            raise serializers.ValidationError(
                f"'{value.name}' is password-protected and cannot be processed."
            )

        page_count = len(document)
    finally:
        document.close()

    if page_count == 0:
        raise serializers.ValidationError(f"'{value.name}' has no pages.")

    if page_count > max_pages:
        raise serializers.ValidationError(
            f"'{value.name}' has {page_count} pages, exceeding the maximum "
            f"of {max_pages}."
        )

    return page_count


def validate_image_file(value, max_size=DEFAULT_MAX_IMAGE_SIZE):
    """
    Validate an uploaded file is a real, decodable JPEG/PNG image (used by
    JPG-to-PDF and anything else that accepts raster images). Never trust
    the extension/Content-Type alone - Pillow actually decodes the pixel
    data, which a renamed non-image file cannot survive.

    Raises rest_framework.serializers.ValidationError with a clear,
    per-file message on any failure. Returns the image's (width, height).
    """
    if value.size > max_size:
        raise serializers.ValidationError(
            f"'{value.name}' exceeds the maximum allowed size of "
            f"{max_size // (1024 * 1024)} MB."
        )

    if not value.name.lower().endswith(ALLOWED_IMAGE_EXTENSIONS):
        raise serializers.ValidationError(
            f"'{value.name}' is not a supported image file (JPG/PNG only)."
        )

    value.seek(0)
    try:
        image = Image.open(value)
        image.verify()
    except Exception:
        raise serializers.ValidationError(
            f"'{value.name}' could not be opened - it may be corrupted or not a real image."
        )
    finally:
        value.seek(0)

    # verify() closes the file object internally - reopen for a real decode
    # (verify() alone doesn't fully decode pixel data) and to read the size.
    try:
        image = Image.open(value)
        image.load()
        width, height = image.size
        image_format = image.format
    except Exception:
        raise serializers.ValidationError(
            f"'{value.name}' could not be decoded - it may be corrupted."
        )
    finally:
        value.seek(0)

    if image_format not in ALLOWED_IMAGE_FORMATS:
        raise serializers.ValidationError(
            f"'{value.name}' is a {image_format} file - only JPG/PNG are supported."
        )

    return width, height


_OFFICE_OPENERS = {}


def _get_office_openers():
    """Lazily imported: python-docx/openpyxl/python-pptx are only needed
    by the X-to-PDF conversion features, not every request that touches
    apps.common.validation."""
    if not _OFFICE_OPENERS:
        from docx import Document as DocxDocument
        from openpyxl import load_workbook
        from pptx import Presentation

        _OFFICE_OPENERS.update({
            "docx": lambda f: DocxDocument(f),
            "xlsx": lambda f: load_workbook(f),
            "pptx": lambda f: Presentation(f),
        })
    return _OFFICE_OPENERS


def validate_office_file(value, kind, max_size=DEFAULT_MAX_OFFICE_FILE_SIZE):
    """
    Validate an uploaded file is a real, openable Office document of the
    given `kind` ("docx", "xlsx", or "pptx") - used by Word/Excel/
    PowerPoint-to-PDF. Never trust the extension/Content-Type alone - this
    actually opens the file with the matching python-docx/openpyxl/
    python-pptx library, which a renamed non-office file cannot survive.

    Raises rest_framework.serializers.ValidationError with a clear
    message on any failure.
    """
    if kind not in ("docx", "xlsx", "pptx"):
        raise ValueError(f"kind must be one of docx/xlsx/pptx (got {kind!r}).")

    if value.size > max_size:
        raise serializers.ValidationError(
            f"'{value.name}' exceeds the maximum allowed size of "
            f"{max_size // (1024 * 1024)} MB."
        )

    if not value.name.lower().endswith(f".{kind}"):
        raise serializers.ValidationError(f"'{value.name}' is not a .{kind} file.")

    value.seek(0)
    header = value.read(len(ZIP_MAGIC_BYTES))
    value.seek(0)
    if header != ZIP_MAGIC_BYTES:
        raise serializers.ValidationError(f"'{value.name}' does not look like a valid .{kind} file.")

    try:
        _get_office_openers()[kind](value)
    except Exception:
        raise serializers.ValidationError(
            f"'{value.name}' could not be opened - it may be corrupted."
        )
    finally:
        value.seek(0)
